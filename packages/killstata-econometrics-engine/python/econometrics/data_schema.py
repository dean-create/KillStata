"""Canonical-table schema normalization.

The importer must not turn identifiers such as ``00123`` into ``123`` merely
because a parser found something numeric-looking.  This module therefore uses
a deliberately narrow conversion rule: only complete columns of plain numeric
text are converted.  Anything carrying a leading zero, date-like text,
percentage sign, thousands separator, or missing-like literal stays textual
until a later explicit transformation decides otherwise.
"""

from __future__ import annotations

import re
from typing import Any

import pandas as pd


_LEADING_ZERO_IDENTIFIER = re.compile(r"^[+-]?0\d+$")
_PLAIN_NUMBER = re.compile(r"^[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?$")
_SPECIAL_NUMERIC = {"inf", "+inf", "-inf", "infinity", "+infinity", "-infinity"}
_DATE_LIKE = re.compile(r"^\d{4}[-/]\d{1,2}(?:[-/]\d{1,2})?$")
_AMBIGUOUS_MISSING = {"na", "n/a", "null", "none", "-", "--"}
_EXCEL_ZERO_NUMBER_FORMAT = re.compile(r"^0+$")


def _textual(series: pd.Series) -> pd.Series:
    """Preserve literal tokens while representing only truly blank cells as missing."""
    values = series.astype("string").str.strip()
    return values.mask(values == "", pd.NA)


def _has_ambiguous_text(values: pd.Series) -> bool:
    for value in values.tolist():
        lowered = value.lower()
        if lowered in _AMBIGUOUS_MISSING:
            return True
        if "%" in value or "," in value or _DATE_LIKE.fullmatch(value):
            return True
    return False


def _is_numeric_literal(value: str) -> bool:
    # DID2S 用 -inf 表示从未处理组；它是方法契约要求的数值哨兵，不能在
    # 导入时退化成类别。其他无穷值仍由后续质量检查决定是否允许进入模型。
    return bool(_PLAIN_NUMBER.fullmatch(value)) or value.lower() in _SPECIAL_NUMERIC


def _column_receipt(
    *, name: str, series: pd.Series, logical_type: str, decision: str, warnings: list[str] | None = None
) -> dict[str, Any]:
    return {
        "name": str(name),
        "physicalType": str(series.dtype),
        "logicalType": logical_type,
        "decision": decision,
        "warnings": warnings or [],
    }


def _restore_excel_zero_padded_identifiers(
    frame: pd.DataFrame,
    *,
    source_path: str,
    source_format: str,
    sheet_policy: dict[str, Any],
) -> tuple[pd.DataFrame, set[str], list[str]]:
    """Restore identifiers represented as numeric cells with a pure ``00000`` format.

    Pandas returns Excel's stored numeric value and intentionally discards the cell
    number format.  For a pure zero-padding format, that format is semantic data:
    ``123`` displayed as ``00123`` is almost always an identifier.  We restore only
    this narrow, reversible case; currency, dates, and arbitrary custom formats stay
    untouched because their display format cannot safely determine logical type.
    """
    if source_format != "xlsx":
        return frame, set(), []

    try:
        from openpyxl import load_workbook

        workbook = load_workbook(source_path, read_only=True, data_only=True)
        mode = sheet_policy.get("mode") or "first_sheet"
        worksheet = workbook[str(sheet_policy["sheetName"])] if mode == "named_sheet" else workbook.worksheets[0]
        header_row = int(sheet_policy.get("headerRow", 0)) + 1
        restored: set[str] = set()
        formatted_by_column: dict[int, dict[int, str]] = {}

        # read_only 工作簿不能逐格调用 worksheet.cell(row, column)：每次随机访问都可能
        # 从压缩流开头重新扫描，真实 did.xlsx 会从不足 1 秒的读取退化到数分钟。
        # 按行顺序流式扫描一次，复杂度稳定为 O(行数×列数)。
        rows = worksheet.iter_rows(
            min_row=header_row + 1,
            max_row=header_row + len(frame),
            max_col=len(frame.columns),
        )
        for dataframe_row, row in enumerate(rows):
            for column_index, cell in enumerate(row, start=1):
                value = cell.value
                number_format = str(cell.number_format or "").split(";", 1)[0].strip()
                if not _EXCEL_ZERO_NUMBER_FORMAT.fullmatch(number_format):
                    continue
                if not isinstance(value, (int, float)) or isinstance(value, bool):
                    continue
                if isinstance(value, float) and not value.is_integer():
                    continue
                formatted_by_column.setdefault(column_index, {})[dataframe_row] = (
                    f"{int(value):0{len(number_format)}d}"
                )

        for column_index, column_name in enumerate(frame.columns, start=1):
            formatted = formatted_by_column.get(column_index, {})
            if formatted:
                values = frame[column_name].astype(object).copy()
                for dataframe_row, value in formatted.items():
                    values.iloc[dataframe_row] = value
                frame[column_name] = values
                restored.add(str(column_name))
        workbook.close()
        return frame, restored, []
    except Exception as exc:
        # Excel 仍可由 pandas 正常导入；无法读取格式时只放弃这项保守恢复，不改变
        # 原始数值或阻断用户导入。receipt 留下 warning 供后续人工判断。
        return frame, set(), [f"Excel zero-padding format recovery skipped: {exc}"]


def normalize_for_canonical(
    frame: pd.DataFrame,
    *,
    source_path: str,
    source_format: str,
    sheet_policy: dict[str, Any],
) -> tuple[pd.DataFrame, dict[str, Any]]:
    """Return a safe canonical frame and a JSON-serializable normalization receipt.

    ``source_path`` and ``sheet_policy`` are part of the public contract even
    when no source-specific correction is needed yet: a receipt must explain
    exactly which import view produced a canonical stage.
    """
    normalized = frame.copy()
    columns: list[dict[str, Any]] = []
    warnings: list[str] = []
    normalized, restored_excel_identifiers, excel_warnings = _restore_excel_zero_padded_identifiers(
        normalized,
        source_path=source_path,
        source_format=source_format,
        sheet_policy=sheet_policy,
    )
    warnings.extend(excel_warnings)

    for name in normalized.columns:
        original = normalized[name]
        if str(name) in restored_excel_identifiers:
            text = _textual(original)
            normalized[name] = text
            columns.append(
                _column_receipt(
                    name=name,
                    series=text,
                    logical_type="identifier",
                    decision="restore_excel_zero_padded_identifier",
                )
            )
            continue
        # Values already provided by typed sources (Stata/Parquet or an Excel
        # numeric cell) remain numeric. Text parsing rules below only decide
        # whether text may safely become numeric.
        if pd.api.types.is_numeric_dtype(original) or pd.api.types.is_datetime64_any_dtype(original):
            logical_type = "number" if pd.api.types.is_numeric_dtype(original) else "datetime"
            columns.append(
                _column_receipt(name=name, series=original, logical_type=logical_type, decision="source_typed")
            )
            continue

        text = _textual(original)
        non_missing = text.dropna()
        if non_missing.empty:
            normalized[name] = text
            columns.append(_column_receipt(name=name, series=text, logical_type="text", decision="preserve_empty_text"))
            continue

        if non_missing.map(lambda value: bool(_LEADING_ZERO_IDENTIFIER.fullmatch(value))).any():
            normalized[name] = text
            columns.append(
                _column_receipt(name=name, series=text, logical_type="identifier", decision="preserve_identifier")
            )
            continue

        if _has_ambiguous_text(non_missing):
            normalized[name] = text
            warning = "ambiguous textual values were preserved; use an explicit later transformation to coerce them"
            columns.append(
                _column_receipt(
                    name=name,
                    series=text,
                    logical_type="text",
                    decision="preserve_ambiguous_text",
                    warnings=[warning],
                )
            )
            warnings.append(f"{name}: {warning}")
            continue

        if non_missing.map(_is_numeric_literal).all():
            normalized[name] = pd.to_numeric(text, errors="raise")
            columns.append(
                _column_receipt(name=name, series=normalized[name], logical_type="number", decision="safe_numeric")
            )
            continue

        normalized[name] = text
        columns.append(_column_receipt(name=name, series=text, logical_type="text", decision="preserve_text"))

    return normalized, {
        "version": 1,
        "sourcePath": source_path,
        "sourceFormat": source_format,
        "sheetPolicy": sheet_policy,
        "columns": columns,
        "warnings": warnings,
    }
